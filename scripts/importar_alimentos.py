#!/usr/bin/env python3
"""
importar_alimentos.py — gera foods.json a partir dos arquivos em dados/.

Reexecutável e idempotente: sempre reconstrói foods.json do zero.
Fontes detectadas automaticamente:
  * TACO 4ª ed. (NEPA/Unicamp) em PDF  -> taco*.pdf (Tabela 1, centesimal + minerais)
  * Planilhas .csv/.xlsx de TACO ou TBCA -> detectadas pelo cabeçalho
    (TBCA tem prioridade em duplicatas, por ser mais recente)

Limpeza: NA / * / - / vazio = ausente (null); "Tr" (traço) = 0, anotado em "tr".
Uso:  python scripts/importar_alimentos.py
Requer: pdfplumber (pip install pdfplumber); openpyxl só se houver .xlsx.
"""
import csv, glob, json, os, re, sys, unicodedata

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DADOS = os.path.join(RAIZ, 'dados')
SAIDA = os.path.join(RAIZ, 'foods.json')
CAMPOS = ['kcal', 'prot', 'carb', 'gord', 'fibra', 'sodio_mg']

# ---------- utilidades ----------

def normalizar(s):
    s = unicodedata.normalize('NFD', s.lower())
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    return re.sub(r'[^a-z0-9]+', ' ', s).strip()

def valor(tok):
    """Converte célula -> (numero|None, traco:bool)."""
    if tok is None:
        return None, False
    t = str(tok).strip()
    t = re.sub(r'(?<=\d)[a-z]+$', '', t)          # remove notas de rodapé (ex.: "20a")
    if t.lower() == 'tr':
        return 0.0, True
    if t in ('', 'NA', 'na', '*', '-', '–', '—'):
        return None, False
    try:
        return float(t.replace(',', '.')), False
    except ValueError:
        return None, False

def item(fid, nome, grupo, fonte, vals):
    """vals: dict campo -> célula bruta."""
    d = {'id': fid, 'nome': nome.strip(), 'grupo': grupo, 'fonte': fonte}
    tr, falta = [], []
    for c in CAMPOS:
        v, t = valor(vals.get(c))
        d[c] = None if v is None else round(v, 2)
        if t: tr.append(c)
        if v is None: falta.append(c)
    if tr: d['tr'] = tr
    if falta: d['falta'] = falta
    return d

# ---------- TACO em PDF ----------
# Tabela 1 alterna páginas: (a) centesimal: nº, descrição, umidade, kcal, kJ, prot, lip,
# colest, carb, fibra, cinzas, Ca, Mg  (11 valores no fim da linha);
# (b) minerais/vitaminas: nº, Mn, P, Fe, Na, ... (células vazias somem no texto,
# por isso o sódio é localizado pela posição x do cabeçalho "Sódio").
NUM = r'(?:\d+(?:,\d+)?[a-z]?|NA|Tr|\*)'
LINHA_A = re.compile(r'^(\d+)\s+(.+?)\s+((?:' + NUM + r'\s+){10}' + NUM + r')$')

# Micronutrientes (por 100 g) guardados em "mic": página (a) = colesterol, Ca, Mg; página (b) = demais.
# Página (b), 15 colunas: Mn, P, Fe, Na, K, Cu, Zn, Retinol, RE, RAE, Tiamina, Riboflavina, Piridoxina, Niacina, Vit. C
MIC_A = {'colest_mg': 5, 'calcio_mg': 9, 'magnesio_mg': 10}
MIC_B = {'fosforo_mg': 1, 'ferro_mg': 2, 'potassio_mg': 4, 'zinco_mg': 6, 'vita_ug': 9, 'tiamina_mg': 10,
         'riboflavina_mg': 11, 'piridoxina_mg': 12, 'niacina_mg': 13, 'vitc_mg': 14}

def micros(va, vb):
    m = {}
    for campo, j in MIC_A.items():
        v, _ = valor(va[j] if va else None)
        if v is not None: m[campo] = round(v, 3)
    for campo, j in MIC_B.items():
        v, _ = valor(vb[j] if vb else None)
        if v is None and campo == 'vita_ug':          # origem animal: só retinol (RAE = retinol)
            v, _ = valor(vb[7] if vb else None)
        if v is not None: m[campo] = round(v, 3)
    return m

def importar_taco_pdf(caminho):
    import pdfplumber
    centesimal, sodio, grupos, minerais = {}, {}, {}, {}
    grupo = ''
    pendente = ''
    with pdfplumber.open(caminho) as pdf:
        for pg in pdf.pages:
            texto = pg.extract_text() or ''
            if not texto.startswith('Tabela 1'):
                continue
            if 'Carbo-' in texto:                       # página (a)
                palavras = pg.extract_words(extra_attrs=['fontname'])
                # linha de unidades "(%) (kcal) (kJ) (g) ..." dá o centro x das 11 colunas
                unid = [w for w in palavras if re.fullmatch(r'\((%|kcal|kJ|g|mg)\)', w['text'])]
                if len(unid) < 11: continue
                unid.sort(key=lambda w: w['x0'])
                cx = [(w['x0'] + w['x1']) / 2 for w in unid]
                topo, borda = unid[0]['bottom'] + 2, unid[0]['x0'] - 8
                linhas = {}
                for w in palavras:
                    if w['top'] > topo:
                        linhas.setdefault(round(w['top']), []).append(w)
                ultimo = None
                for t in sorted(linhas):
                    ws = sorted(linhas[t], key=lambda w: w['x0'])
                    desc = [w for w in ws if w['x1'] < borda]
                    vals = [w for w in ws if w['x1'] >= borda]
                    if not desc: continue
                    if re.fullmatch(r'\d+', desc[0]['text']) and vals:
                        n = int(desc[0]['text'])
                        v = [None] * 11
                        for w in vals:
                            j = min(range(11), key=lambda k: abs((w['x0'] + w['x1']) / 2 - cx[k]))
                            v[j] = w['text']
                        centesimal[n] = (' '.join(w['text'] for w in desc[1:]), v)
                        grupos[n] = grupo; ultimo = n
                    elif 'Bold' in desc[0]['fontname']:
                        grupo = ' '.join(w['text'] for w in desc); ultimo = None
                    elif ultimo and not vals and not re.fullmatch(r'\d+', desc[0]['text']):
                        nome, v = centesimal[ultimo]           # descrição quebrada em 2 linhas
                        centesimal[ultimo] = (nome + ' ' + ' '.join(w['text'] for w in desc), v)
            elif 'Sódio' in texto:                          # página (b)
                palavras = pg.extract_words()
                cab = [w for w in palavras if w['text'] == 'Sódio']
                if not cab: continue
                xs = (cab[0]['x0'] + cab[0]['x1']) / 2
                topo = cab[0]['bottom'] + 12
                # linha de unidades "(mg) (µg)" logo abaixo do cabeçalho: centro x das 15 colunas
                unid = sorted([w for w in palavras if re.fullmatch(r'\((mg|µg|mcg)\)', w['text'])
                               and cab[0]['bottom'] - 2 < w['top'] < cab[0]['bottom'] + 14], key=lambda w: w['x0'])
                cols = [(w['x0'] + w['x1']) / 2 for w in unid] if len(unid) == 15 else None
                linhas = {}
                for w in palavras:
                    if w['top'] > topo:
                        linhas.setdefault(round(w['top']), []).append(w)
                for ws in linhas.values():
                    ws.sort(key=lambda w: w['x0'])
                    if not re.fullmatch(r'\d+', ws[0]['text']) or len(ws) < 3:
                        continue
                    n = int(ws[0]['text'])
                    alvo = min(ws[1:], key=lambda w: abs((w['x0'] + w['x1']) / 2 - xs))
                    if abs((alvo['x0'] + alvo['x1']) / 2 - xs) < 14:
                        sodio[n] = alvo['text']
                    if cols:
                        v = [None] * 15
                        for w in ws[1:]:
                            c = (w['x0'] + w['x1']) / 2
                            j = min(range(15), key=lambda k: abs(c - cols[k]))
                            if abs(c - cols[j]) < 14: v[j] = w['text']
                        minerais[n] = v
    itens = []
    for n in sorted(centesimal):
        nome, v = centesimal[n]
        vals = {'kcal': v[1], 'prot': v[3], 'gord': v[4], 'carb': v[6], 'fibra': v[7],
                'sodio_mg': sodio.get(n)}
        it = item(f'taco-{n}', nome, grupos[n], 'TACO', vals)
        mic = micros(v, minerais.get(n))
        if mic: it['mic'] = mic
        itens.append(it)
    print(f'  micronutrientes: {sum(1 for i in itens if "mic" in i)} alimentos com ao menos um valor')
    return itens

# ---------- planilhas (CSV/XLSX) de TACO ou TBCA ----------
CHAVES = {  # campo -> palavras que identificam a coluna no cabeçalho normalizado
    'nome': ['descricao', 'alimento', 'nome'],
    'kcal': ['kcal', 'energia'], 'prot': ['proteina'], 'gord': ['lipideo', 'lipidio', 'gordura'],
    'carb': ['carboidrato'], 'fibra': ['fibra'], 'sodio_mg': ['sodio'],
}

def ler_planilha(caminho):
    if caminho.lower().endswith('.csv'):
        with open(caminho, encoding='utf-8-sig', errors='replace') as f:
            amostra = f.read(4096); f.seek(0)
            sep = ';' if amostra.count(';') > amostra.count(',') else ','
            return [r for r in csv.reader(f, delimiter=sep)]
    import openpyxl
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    linhas = []
    for ws in wb.worksheets:
        linhas += [[('' if c is None else str(c)) for c in r] for r in ws.iter_rows(values_only=True)]
    return linhas

def importar_planilha(caminho):
    linhas = ler_planilha(caminho)
    # cabeçalho pode ocupar várias linhas (células mescladas): junta até 4 linhas por coluna
    for i in range(min(30, len(linhas))):
        bloco = linhas[i:i + 4]
        largura = max(len(r) for r in bloco)
        cab = [normalizar(' '.join(r[j] if j < len(r) else '' for r in bloco)) for j in range(largura)]
        col = {}
        for campo, chaves in CHAVES.items():
            for j, h in enumerate(cab):
                if any(k in h for k in chaves) and j not in col.values():
                    if campo == 'kcal' and 'kj' in h and 'kcal' not in h: continue
                    col[campo] = j; break
        if {'nome', 'kcal', 'prot', 'carb'} <= col.keys():
            break
    else:
        print(f'  ! cabeçalho não reconhecido: {os.path.basename(caminho)}'); return [], None
    texto = ' '.join(' '.join(r) for r in linhas[:10]).lower()
    fonte = 'TBCA' if 'tbca' in texto or 'tbca' in caminho.lower() else 'TACO'
    itens = []
    for k, r in enumerate(linhas[i + 4:]):
        nome = r[col['nome']].strip() if col['nome'] < len(r) else ''
        if not nome or valor(r[col['kcal']] if col['kcal'] < len(r) else None)[0] is None:
            continue
        vals = {c: (r[j] if j < len(r) else None) for c, j in col.items() if c != 'nome'}
        itens.append(item(f'{fonte.lower()}-p{k}', nome, '', fonte, vals))
    return itens, fonte

# ---------- principal ----------

def main():
    arquivos = sorted(glob.glob(os.path.join(DADOS, '*')))
    por_fonte = {'TBCA': [], 'TACO': []}
    for a in arquivos:
        nome = os.path.basename(a).lower()
        if nome.endswith('.pdf') and 'taco' in nome:
            itens = importar_taco_pdf(a); por_fonte['TACO'] += itens
            print(f'  TACO (PDF) {os.path.basename(a)}: {len(itens)}')
        elif nome.endswith(('.csv', '.xlsx')):
            itens, fonte = importar_planilha(a)
            if fonte:
                por_fonte[fonte] += itens; print(f'  {fonte} {os.path.basename(a)}: {len(itens)}')
        else:
            print(f'  (ignorado) {os.path.basename(a)}')
    vistos, base, dup = set(), [], 0
    for fonte in ('TBCA', 'TACO'):                      # TBCA primeiro: vence em conflito
        for it in por_fonte[fonte]:
            k = normalizar(it['nome'])
            if k in vistos: dup += 1; continue
            vistos.add(k); base.append(it)
    with open(SAIDA, 'w', encoding='utf-8') as f:
        json.dump(base, f, ensure_ascii=False, separators=(',', ':'))
    # relatório
    falta = [it for it in base if 'falta' in it]
    incoer = [it for it in base if None not in (it['kcal'], it['prot'], it['carb'], it['gord'])
              and abs(it['kcal'] - (4 * it['prot'] + 4 * it['carb'] + 9 * it['gord'])) > max(40, 0.25 * it['kcal'])]
    print(f'Total importado: {len(base)} | duplicatas removidas: {dup} | com campos faltando: {len(falta)}')
    for c in CAMPOS:
        print(f'  sem {c}: {sum(1 for it in base if it[c] is None)}')
    print(f'  kcal muito diferente de 4P+4C+9G (conferir): {len(incoer)}')
    if not base:
        sys.exit(1)

if __name__ == '__main__':
    main()
